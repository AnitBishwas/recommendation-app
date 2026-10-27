from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# ── Brand palette ──────────────────────────────────────────────────────────────
DARK_BG    = RGBColor(0x0F, 0x17, 0x2A)
ACCENT     = RGBColor(0x00, 0xB5, 0x8A)
ACCENT2    = RGBColor(0xF7, 0xC9, 0x48)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xB0, 0xBE, 0xC5)
CARD_BG    = RGBColor(0x16, 0x22, 0x3A)

# Mockup UI colours
UI_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
UI_BG      = RGBColor(0xF6, 0xF6, 0xF7)   # Polaris page bg
UI_BORDER  = RGBColor(0xD9, 0xDB, 0xDE)
UI_TEXT    = RGBColor(0x1A, 0x1A, 0x1A)
UI_MUTED   = RGBColor(0x6D, 0x7A, 0x8C)
UI_GREEN   = RGBColor(0x00, 0x8C, 0x5E)   # Polaris success
UI_GREEN_BG= RGBColor(0xD4, 0xED, 0xE4)
UI_RED_BG  = RGBColor(0xFD, 0xE8, 0xE8)
UI_RED     = RGBColor(0xBF, 0x0E, 0x08)
UI_HEADER  = RGBColor(0x1A, 0x1A, 0x1A)   # Polaris nav bg
UI_NAV_TXT = RGBColor(0xF1, 0xF1, 0xF1)
UI_INPUT   = RGBColor(0xF9, 0xF9, 0xF9)
UI_ACCENT  = RGBColor(0x00, 0x8C, 0x5E)   # Polaris primary button

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]

TOTAL = 13   # updated count

# ── Helpers ────────────────────────────────────────────────────────────────────

def add_rect(slide, x, y, w, h, fill_rgb, radius=False):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.fill.background()
    shape.line.width = 0
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    return shape


def add_rect_border(slide, x, y, w, h, fill_rgb, border_rgb, border_pt=1):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.color.rgb = border_rgb
    shape.line.width = Pt(border_pt)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    return shape


def txbox(slide, text, x, y, w, h,
          size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
          wrap=True, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def bg(slide):
    add_rect(slide, 0, 0, W, H, DARK_BG)


def accent_bar(slide, y=Inches(0.55), h=Inches(0.06)):
    add_rect(slide, 0, y, W, h, ACCENT)


def slide_number(slide, n, total):
    txbox(slide, f"{n} / {total}", Inches(12.0), Inches(7.1),
          Inches(1.2), Inches(0.35), size=11, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)


# ══════════════════════════════════════════════════════════════════════════════
# WIREFRAME HELPER COMPONENTS
# ══════════════════════════════════════════════════════════════════════════════

def ui_polaris_nav(slide, x, y, w):
    """Shopify-style top nav bar."""
    h = Inches(0.38)
    add_rect(slide, x, y, w, h, UI_HEADER)
    txbox(slide, "Shopify", x + Inches(0.15), y + Inches(0.05),
          Inches(0.9), Inches(0.28), size=11, bold=True, color=UI_NAV_TXT)
    txbox(slide, "Recommendation Engine App",
          x + Inches(1.1), y + Inches(0.05),
          Inches(3), Inches(0.28), size=9, color=LIGHT_GREY)
    return h


def ui_badge(slide, x, y, label, active=True):
    col = UI_GREEN_BG if active else UI_RED_BG
    txt_col = UI_GREEN if active else UI_RED
    add_rect_border(slide, x, y, Inches(0.7), Inches(0.22), col, col, 0)
    txbox(slide, label, x, y, Inches(0.7), Inches(0.22),
          size=7, bold=True, color=txt_col, align=PP_ALIGN.CENTER)


def ui_btn_primary(slide, x, y, label, w=Inches(1.3), h=Inches(0.28)):
    add_rect(slide, x, y, w, h, UI_ACCENT)
    txbox(slide, label, x, y, w, h, size=9, bold=True,
          color=UI_WHITE, align=PP_ALIGN.CENTER)


def ui_btn_ghost(slide, x, y, label, w=Inches(1.1), h=Inches(0.28)):
    add_rect_border(slide, x, y, w, h, UI_WHITE, UI_BORDER, 1)
    txbox(slide, label, x, y, w, h, size=9, color=UI_TEXT,
          align=PP_ALIGN.CENTER)


def ui_input(slide, x, y, w, label, value="", h=Inches(0.32)):
    add_rect_border(slide, x, y, w, h, UI_INPUT, UI_BORDER, 1)
    if value:
        txbox(slide, value, x + Inches(0.08), y + Inches(0.04),
              w - Inches(0.12), h - Inches(0.06), size=9, color=UI_TEXT)
    else:
        txbox(slide, label, x + Inches(0.08), y + Inches(0.04),
              w - Inches(0.12), h - Inches(0.06), size=9, color=UI_MUTED, italic=True)


def ui_label(slide, x, y, text, w=Inches(2)):
    txbox(slide, text, x, y, w, Inches(0.2), size=8, bold=True, color=UI_TEXT)


def ui_card(slide, x, y, w, h):
    add_rect_border(slide, x, y, w, h, UI_WHITE, UI_BORDER, 1)
    return (x, y, w, h)


def ui_section_header(slide, x, y, w, title):
    add_rect(slide, x, y, w, Inches(0.35), UI_WHITE)
    txbox(slide, title, x + Inches(0.15), y + Inches(0.06),
          w - Inches(0.2), Inches(0.25), size=11, bold=True, color=UI_HEADER)
    add_rect(slide, x, y + Inches(0.35), w, Pt(1), UI_BORDER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDES
# ══════════════════════════════════════════════════════════════════════════════

slides_data = []


# ── 1. Title ──────────────────────────────────────────────────────────────────
def slide_title(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    add_rect(sl, 0, Inches(2.8), W, Inches(2.1), CARD_BG)
    accent_bar(sl, y=Inches(2.75), h=Inches(0.06))
    accent_bar(sl, y=Inches(4.85), h=Inches(0.06))
    txbox(sl, "Recommendation Engine", Inches(1.0), Inches(2.9),
          Inches(11.3), Inches(0.9), size=46, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txbox(sl, "Shopify · Personalised Product Discovery via UTM-Targeted Quizzes",
          Inches(1.0), Inches(3.85), Inches(11.3), Inches(0.5),
          size=18, color=ACCENT, align=PP_ALIGN.CENTER, italic=True)
    txbox(sl, "Built on Node.js · Express · MongoDB · Shopify Admin API",
          Inches(1.0), Inches(5.5), Inches(11.3), Inches(0.4),
          size=13, color=LIGHT_GREY, align=PP_ALIGN.CENTER)
    slide_number(sl, n, total)

slides_data.append(slide_title)


# ── 2. Problem Statement ───────────────────────────────────────────────────────
def slide_problem(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "The Problem", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    cards = [
        ("Generic storefronts", "Visitors land from specific campaigns but see the same generic catalog — no personalisation at all."),
        ("Discovery friction", "Customers don't know which product fits their needs, leading to high bounce rates and low conversion."),
        ("No-code gap", "Merchants can't wire marketing UTM data to product recommendations without custom dev work."),
        ("Lost revenue", "Missed up-sell opportunities — the right product exists in the catalog but never reaches the right shopper."),
    ]
    col_w = Inches(3.0)
    col_gap = Inches(0.18)
    start_x = Inches(0.55)
    card_y = Inches(1.55)
    card_h = Inches(4.8)
    for i, (title, body) in enumerate(cards):
        x = start_x + i * (col_w + col_gap)
        add_rect(sl, x, card_y, col_w, card_h, CARD_BG)
        add_rect(sl, x, card_y, col_w, Inches(0.07), ACCENT2)
        txbox(sl, title, x + Inches(0.18), card_y + Inches(0.18),
              col_w - Inches(0.3), Inches(0.55), size=15, bold=True, color=ACCENT2)
        txbox(sl, body, x + Inches(0.18), card_y + Inches(0.75),
              col_w - Inches(0.3), Inches(3.8), size=13, color=LIGHT_GREY, wrap=True)
    slide_number(sl, n, total)

slides_data.append(slide_problem)


# ── 3. Solution Overview ───────────────────────────────────────────────────────
def slide_solution(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "The Solution", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)
    txbox(sl, "A Shopify embedded app that turns UTM campaign data into personalised quiz-driven product recommendations — zero coding required.",
          Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.7), size=16, color=LIGHT_GREY, wrap=True)

    flow = [
        ("UTM Link", "Visitor clicks a\ncampaign link with\nutm_source / medium /\ncampaign params"),
        ("Quiz Trigger", "App Proxy matches\nUTM params to an\nactive campaign and\nserves the quiz"),
        ("Scored Answers", "Each answer maps to\nproduct tags + weights.\nScores accumulate\nacross questions"),
        ("Recommendations", "Top-scored products\nfetched from Shopify\ncatalog and ranked\nfor the visitor"),
    ]
    box_w = Inches(2.8)
    box_h = Inches(3.2)
    gap   = Inches(0.45)
    start_x = Inches(0.55)
    box_y = Inches(2.3)
    arrow_y = box_y + Inches(1.4)
    for i, (title, desc) in enumerate(flow):
        x = start_x + i * (box_w + gap)
        add_rect(sl, x, box_y, box_w, box_h, CARD_BG)
        add_rect(sl, x, box_y, box_w, Inches(0.07), ACCENT)
        num_x = x + Inches(0.15)
        num_y = box_y + Inches(0.15)
        add_rect(sl, num_x, num_y, Inches(0.38), Inches(0.38), ACCENT)
        txbox(sl, str(i + 1), num_x, num_y - Inches(0.02), Inches(0.38), Inches(0.4),
              size=13, bold=True, color=DARK_BG, align=PP_ALIGN.CENTER)
        txbox(sl, title, x + Inches(0.18), box_y + Inches(0.65),
              box_w - Inches(0.3), Inches(0.45), size=15, bold=True, color=WHITE)
        txbox(sl, desc, x + Inches(0.18), box_y + Inches(1.15),
              box_w - Inches(0.3), Inches(1.9), size=12, color=LIGHT_GREY, wrap=True)
        if i < len(flow) - 1:
            ax = x + box_w + Inches(0.05)
            txbox(sl, "→", ax, arrow_y, gap + Inches(0.05), Inches(0.4),
                  size=22, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    slide_number(sl, n, total)

slides_data.append(slide_solution)


# ── 4. Admin UI Mockup ────────────────────────────────────────────────────────
def slide_admin_mockup(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Admin UI — Campaign Manager & Quiz Builder", Inches(0.5), Inches(0.7),
          Inches(12.3), Inches(0.55), size=28, bold=True, color=WHITE)

    # ── LEFT PANEL: Campaign List ──────────────────────────────────────────
    px = Inches(0.35)
    py = Inches(1.45)
    pw = Inches(5.8)
    ph = Inches(5.75)
    add_rect(sl, px, py, pw, ph, UI_BG)

    nav_h = ui_polaris_nav(sl, px, py, pw)
    # Page header
    header_y = py + nav_h + Inches(0.15)
    add_rect(sl, px, header_y, pw, Inches(0.42), UI_WHITE)
    txbox(sl, "Campaigns", px + Inches(0.18), header_y + Inches(0.08),
          Inches(2), Inches(0.28), size=14, bold=True, color=UI_HEADER)
    ui_btn_primary(sl, px + pw - Inches(1.48), header_y + Inches(0.07),
                   "+ Create campaign", w=Inches(1.38))

    # Table header
    tbl_y = header_y + Inches(0.42) + Inches(0.1)
    add_rect(sl, px + Inches(0.12), tbl_y, pw - Inches(0.24), Inches(0.28), UI_BG)
    cols = [("Campaign Name", 1.65), ("UTM Source", 0.95), ("UTM Medium", 0.95), ("Status", 0.72), ("", 0.6)]
    cx = px + Inches(0.18)
    for (col_lbl, col_w_in) in cols:
        txbox(sl, col_lbl, cx, tbl_y + Inches(0.05), Inches(col_w_in), Inches(0.2),
              size=8, bold=True, color=UI_MUTED)
        cx += Inches(col_w_in)

    # Table rows
    rows = [
        ("Bold Lip Summer", "instagram", "paid_social", True),
        ("Skincare Starter", "email", "newsletter", True),
        ("Foundation Finder", "google", "cpc", False),
        ("Holiday Glow Kit", "tiktok", "organic", True),
        ("Fragrance Fit", "facebook", "retargeting", False),
    ]
    row_h = Inches(0.4)
    for ri, (name, src, medium, active) in enumerate(rows):
        ry = tbl_y + Inches(0.3) + ri * row_h
        row_bg = UI_WHITE if ri % 2 == 0 else RGBColor(0xF9, 0xF9, 0xFA)
        add_rect(sl, px + Inches(0.12), ry, pw - Inches(0.24), row_h, row_bg)
        add_rect(sl, px + Inches(0.12), ry + row_h - Pt(0.5), pw - Inches(0.24), Pt(0.5), UI_BORDER)
        vals = [name, src, medium]
        widths = [1.65, 0.95, 0.95]
        cx2 = px + Inches(0.18)
        for val, ww in zip(vals, widths):
            txbox(sl, val, cx2, ry + Inches(0.1), Inches(ww), Inches(0.22),
                  size=9, color=UI_TEXT)
            cx2 += Inches(ww)
        ui_badge(sl, cx2, ry + Inches(0.1), "Active" if active else "Inactive", active)
        cx2 += Inches(0.75)
        txbox(sl, "Edit", cx2, ry + Inches(0.1), Inches(0.5), Inches(0.22),
              size=9, color=UI_ACCENT, bold=True)

    # label
    txbox(sl, "Campaign List", px + Inches(0.18), py + ph - Inches(0.25),
          Inches(2), Inches(0.22), size=8, italic=True, color=UI_MUTED)

    # ── RIGHT PANEL: Quiz Builder ──────────────────────────────────────────
    rx2 = Inches(6.55)
    ry2 = py
    rw  = Inches(6.44)
    add_rect(sl, rx2, ry2, rw, ph, UI_BG)

    nav_h2 = ui_polaris_nav(sl, rx2, ry2, rw)

    # Page title row
    phdr_y = ry2 + nav_h2 + Inches(0.12)
    add_rect(sl, rx2, phdr_y, rw, Inches(0.38), UI_WHITE)
    txbox(sl, "← Bold Lip Summer — Quiz Builder", rx2 + Inches(0.15), phdr_y + Inches(0.07),
          Inches(3.5), Inches(0.26), size=12, bold=True, color=UI_HEADER)
    ui_btn_primary(sl, rx2 + rw - Inches(1.65), phdr_y + Inches(0.05),
                   "Preview Quiz", w=Inches(1.2))
    ui_btn_ghost(sl, rx2 + rw - Inches(2.88), phdr_y + Inches(0.05), "Activate", w=Inches(1.1))

    # Two-col layout inside
    lcol_x = rx2 + Inches(0.12)
    lcol_w = Inches(2.65)
    rcol_x = lcol_x + lcol_w + Inches(0.18)
    rcol_w = rw - lcol_w - Inches(0.55)
    content_y = phdr_y + Inches(0.42)

    # Left: question list
    ui_card(sl, lcol_x, content_y, lcol_w, Inches(4.75))
    txbox(sl, "Questions", lcol_x + Inches(0.12), content_y + Inches(0.1),
          Inches(1.4), Inches(0.22), size=10, bold=True, color=UI_HEADER)
    questions = [
        ("1", "What's your skin type?", True),
        ("2", "Preferred finish?", False),
        ("3", "Coverage level?", False),
        ("4", "Skin concern?", False),
    ]
    for qi, (qnum, qtxt, selected) in enumerate(questions):
        qy = content_y + Inches(0.4) + qi * Inches(0.88)
        qbg = RGBColor(0xE8, 0xF5, 0xF0) if selected else UI_WHITE
        qborder = UI_GREEN if selected else UI_BORDER
        add_rect_border(sl, lcol_x + Inches(0.1), qy, lcol_w - Inches(0.2), Inches(0.72), qbg, qborder, 1)
        txbox(sl, f"Q{qnum}", lcol_x + Inches(0.2), qy + Inches(0.1), Inches(0.3), Inches(0.2),
              size=7, bold=True, color=UI_GREEN if selected else UI_MUTED)
        txbox(sl, qtxt, lcol_x + Inches(0.22), qy + Inches(0.3), lcol_w - Inches(0.55), Inches(0.32),
              size=9, color=UI_HEADER if selected else UI_TEXT)
        # up/down arrows
        txbox(sl, "↑↓", lcol_x + lcol_w - Inches(0.38), qy + Inches(0.24),
              Inches(0.28), Inches(0.22), size=8, color=UI_MUTED, align=PP_ALIGN.CENTER)
    txbox(sl, "+ Add Question", lcol_x + Inches(0.1), content_y + Inches(4.0),
          lcol_w - Inches(0.2), Inches(0.28), size=9, color=UI_ACCENT, bold=True, align=PP_ALIGN.CENTER)

    # Right: options + scoring rules
    ui_card(sl, rcol_x, content_y, rcol_w, Inches(4.75))
    txbox(sl, "Q1: What's your skin type?   (single-select)",
          rcol_x + Inches(0.12), content_y + Inches(0.1),
          rcol_w - Inches(0.2), Inches(0.22), size=10, bold=True, color=UI_HEADER)

    options_data = [
        ("Dry", [("hydrating", 8), ("nourishing", 6)]),
        ("Oily", [("matte", 9), ("oil-control", 7)]),
        ("Combo", [("balancing", 7), ("matte", 4)]),
    ]
    for oi, (opt_label, rules) in enumerate(options_data):
        oy = content_y + Inches(0.42) + oi * Inches(1.35)
        # option row
        add_rect_border(sl, rcol_x + Inches(0.1), oy, rcol_w - Inches(0.2), Inches(0.28),
                        RGBColor(0xF0, 0xF9, 0xF5), UI_GREEN, 1)
        txbox(sl, f"Option: {opt_label}", rcol_x + Inches(0.2), oy + Inches(0.05),
              Inches(1.4), Inches(0.2), size=9, bold=True, color=UI_GREEN)
        # scoring rules
        for ri2, (tag, wt) in enumerate(rules):
            rule_y = oy + Inches(0.33) + ri2 * Inches(0.46)
            add_rect(sl, rcol_x + Inches(0.25), rule_y, rcol_w - Inches(0.4), Inches(0.35), UI_BG)
            txbox(sl, "tag:", rcol_x + Inches(0.3), rule_y + Inches(0.07), Inches(0.28), Inches(0.2),
                  size=7, color=UI_MUTED)
            # tag input
            add_rect_border(sl, rcol_x + Inches(0.58), rule_y + Inches(0.05),
                            Inches(1.2), Inches(0.25), UI_WHITE, UI_BORDER, 1)
            txbox(sl, tag, rcol_x + Inches(0.63), rule_y + Inches(0.07), Inches(1.1), Inches(0.2),
                  size=8, color=UI_TEXT)
            txbox(sl, "w:", rcol_x + Inches(1.85), rule_y + Inches(0.07), Inches(0.2), Inches(0.2),
                  size=7, color=UI_MUTED)
            add_rect_border(sl, rcol_x + Inches(2.05), rule_y + Inches(0.05),
                            Inches(0.35), Inches(0.25), UI_WHITE, UI_BORDER, 1)
            txbox(sl, str(wt), rcol_x + Inches(2.08), rule_y + Inches(0.07), Inches(0.3), Inches(0.2),
                  size=8, color=UI_ACCENT, bold=True)

    # label
    txbox(sl, "Quiz Builder", rx2 + Inches(0.18), ry2 + ph - Inches(0.25),
          Inches(2), Inches(0.22), size=8, italic=True, color=UI_MUTED)

    # Panel divider label
    slide_number(sl, n, total)

slides_data.append(slide_admin_mockup)


# ── 5. Admin Preview Mockup ───────────────────────────────────────────────────
def slide_admin_preview(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Admin UI — Campaign Preview & Recommendation Results", Inches(0.5), Inches(0.7),
          Inches(12.3), Inches(0.55), size=28, bold=True, color=WHITE)

    # ── LEFT: Campaign Settings Form ──────────────────────────────────────
    fx = Inches(0.35)
    fy = Inches(1.45)
    fw = Inches(5.8)
    fh = Inches(5.75)
    add_rect(sl, fx, fy, fw, fh, UI_BG)
    ui_polaris_nav(sl, fx, fy, fw)
    form_y = fy + Inches(0.5)
    add_rect(sl, fx, form_y, fw, Inches(0.38), UI_WHITE)
    txbox(sl, "← Campaign Settings — Bold Lip Summer",
          fx + Inches(0.15), form_y + Inches(0.08), Inches(3.8), Inches(0.26),
          size=12, bold=True, color=UI_HEADER)

    card_y = form_y + Inches(0.48)
    add_rect_border(sl, fx + Inches(0.15), card_y, fw - Inches(0.3), Inches(2.5), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "Campaign Details", fx + Inches(0.28), card_y + Inches(0.1), Inches(2), Inches(0.22),
          size=10, bold=True, color=UI_HEADER)

    fields = [
        ("Campaign Name", "Bold Lip Summer", 0.55),
        ("UTM Source",    "instagram",       1.05),
        ("UTM Medium",    "paid_social",     1.55),
        ("UTM Campaign",  "bold_lip_2026",   2.05),
    ]
    for label, val, offset in fields:
        ui_label(sl, fx + Inches(0.28), card_y + Inches(offset), label)
        ui_input(sl, fx + Inches(0.28), card_y + Inches(offset + 0.2),
                 fw - Inches(0.55), label, value=val)

    # Match mode
    mm_y = card_y + Inches(2.6)
    add_rect_border(sl, fx + Inches(0.15), mm_y, fw - Inches(0.3), Inches(0.85), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "Match Mode", fx + Inches(0.28), mm_y + Inches(0.1), Inches(1.5), Inches(0.22),
          size=10, bold=True, color=UI_HEADER)
    # AND pill (selected)
    add_rect(sl, fx + Inches(0.28), mm_y + Inches(0.38), Inches(0.7), Inches(0.26), UI_GREEN)
    txbox(sl, "ALL (AND)", fx + Inches(0.28), mm_y + Inches(0.38), Inches(0.7), Inches(0.26),
          size=8, bold=True, color=UI_WHITE, align=PP_ALIGN.CENTER)
    add_rect_border(sl, fx + Inches(1.05), mm_y + Inches(0.38), Inches(0.7), Inches(0.26),
                    UI_WHITE, UI_BORDER, 1)
    txbox(sl, "ANY (OR)", fx + Inches(1.05), mm_y + Inches(0.38), Inches(0.7), Inches(0.26),
          size=8, color=UI_TEXT, align=PP_ALIGN.CENTER)

    # Result limits
    rl_y = mm_y + Inches(0.95)
    add_rect_border(sl, fx + Inches(0.15), rl_y, fw - Inches(0.3), Inches(1.0), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "Result Limits", fx + Inches(0.28), rl_y + Inches(0.1), Inches(1.5), Inches(0.22),
          size=10, bold=True, color=UI_HEADER)
    ui_label(sl, fx + Inches(0.28), rl_y + Inches(0.35), "Max Tags")
    add_rect_border(sl, fx + Inches(0.28), rl_y + Inches(0.55), Inches(0.55), Inches(0.28),
                    UI_WHITE, UI_BORDER, 1)
    txbox(sl, "5", fx + Inches(0.28), rl_y + Inches(0.55), Inches(0.55), Inches(0.28),
          size=10, bold=True, color=UI_TEXT, align=PP_ALIGN.CENTER)
    ui_label(sl, fx + Inches(1.2), rl_y + Inches(0.35), "Max Products")
    add_rect_border(sl, fx + Inches(1.2), rl_y + Inches(0.55), Inches(0.55), Inches(0.28),
                    UI_WHITE, UI_BORDER, 1)
    txbox(sl, "6", fx + Inches(1.2), rl_y + Inches(0.55), Inches(0.55), Inches(0.28),
          size=10, bold=True, color=UI_TEXT, align=PP_ALIGN.CENTER)

    txbox(sl, "Campaign Settings", fx + Inches(0.18), fy + fh - Inches(0.25),
          Inches(2.5), Inches(0.22), size=8, italic=True, color=UI_MUTED)

    # ── RIGHT: Preview Result Screen ──────────────────────────────────────
    prx = Inches(6.55)
    pry = fy
    prw = Inches(6.44)
    add_rect(sl, prx, pry, prw, fh, UI_BG)
    ui_polaris_nav(sl, prx, pry, prw)
    prev_y = pry + Inches(0.5)
    add_rect(sl, prx, prev_y, prw, Inches(0.38), UI_WHITE)
    txbox(sl, "Preview — Quiz Results", prx + Inches(0.15), prev_y + Inches(0.08),
          Inches(3), Inches(0.26), size=12, bold=True, color=UI_HEADER)

    # Preview quiz answers summary
    ans_y = prev_y + Inches(0.5)
    add_rect_border(sl, prx + Inches(0.15), ans_y, prw - Inches(0.3), Inches(1.2), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "Your Answers", prx + Inches(0.28), ans_y + Inches(0.08), Inches(2), Inches(0.22),
          size=10, bold=True, color=UI_HEADER)
    answers_summary = [
        ("Q1: Skin type", "Dry"),
        ("Q2: Finish", "Matte"),
        ("Q3: Coverage", "Medium"),
    ]
    for ai, (q, a) in enumerate(answers_summary):
        ax2 = prx + Inches(0.28) + ai * Inches(2.0)
        txbox(sl, q, ax2, ans_y + Inches(0.36), Inches(1.5), Inches(0.18),
              size=7, color=UI_MUTED)
        txbox(sl, a, ax2, ans_y + Inches(0.55), Inches(1.5), Inches(0.22),
              size=9, bold=True, color=UI_GREEN)

    # Product cards
    prod_y = ans_y + Inches(1.3)
    add_rect_border(sl, prx + Inches(0.15), prod_y, prw - Inches(0.3), Inches(3.7), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "Recommended Products", prx + Inches(0.28), prod_y + Inches(0.1), Inches(2.5), Inches(0.22),
          size=10, bold=True, color=UI_HEADER)

    products = [
        ("Matte Fix Foundation", "$42.00", 19),
        ("Velvet Lip Colour", "$28.00", 14),
        ("Hydra-Matte Serum", "$56.00", 11),
        ("Long-Wear Concealer", "$35.00", 9),
    ]
    card_w = Inches(1.35)
    card_gap = Inches(0.12)
    for pi, (pname, price, score) in enumerate(products):
        pcx = prx + Inches(0.28) + pi * (card_w + card_gap)
        pcy = prod_y + Inches(0.4)
        # image placeholder
        add_rect_border(sl, pcx, pcy, card_w, Inches(1.1), UI_BG, UI_BORDER, 1)
        txbox(sl, "[ img ]", pcx, pcy, card_w, Inches(1.1),
              size=8, color=UI_MUTED, align=PP_ALIGN.CENTER)
        # info
        txbox(sl, pname, pcx, pcy + Inches(1.15), card_w, Inches(0.28),
              size=8, bold=True, color=UI_TEXT, wrap=True)
        txbox(sl, price, pcx, pcy + Inches(1.43), card_w, Inches(0.2),
              size=8, color=UI_MUTED)
        # score badge
        add_rect(sl, pcx, pcy + Inches(1.65), card_w, Inches(0.22), UI_GREEN_BG)
        txbox(sl, f"Score: {score}", pcx, pcy + Inches(1.65), card_w, Inches(0.22),
              size=7, bold=True, color=UI_GREEN, align=PP_ALIGN.CENTER)
        # ATC
        add_rect(sl, pcx, pcy + Inches(1.9), card_w, Inches(0.26), UI_ACCENT)
        txbox(sl, "Add to Cart", pcx, pcy + Inches(1.9), card_w, Inches(0.26),
              size=7, bold=True, color=UI_WHITE, align=PP_ALIGN.CENTER)

    txbox(sl, "Preview Modal", prx + Inches(0.18), pry + fh - Inches(0.25),
          Inches(2.5), Inches(0.22), size=8, italic=True, color=UI_MUTED)

    slide_number(sl, n, total)

slides_data.append(slide_admin_preview)


# ── 6. Customer Journey Mockup (Storefront) ───────────────────────────────────
def slide_customer_mockup(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Customer UI — Storefront Quiz Flow (recommendation.liquid)", Inches(0.5), Inches(0.7),
          Inches(12.3), Inches(0.55), size=28, bold=True, color=WHITE)

    # Store page bg strip behind modals
    add_rect(sl, Inches(0.3), Inches(1.4), Inches(12.73), Inches(5.85), CARD_BG)
    txbox(sl, "← Shopify Storefront   mystoredemo.myshopify.com/pages/find-your-match?utm_source=instagram&utm_campaign=bold_lip",
          Inches(0.45), Inches(1.48), Inches(12.3), Inches(0.22),
          size=7, color=LIGHT_GREY)

    # ── STATE 1: Intro Screen ──────────────────────────────────────────────
    s1x = Inches(0.55)
    s1y = Inches(1.75)
    s1w = Inches(3.7)
    s1h = Inches(5.1)
    add_rect_border(sl, s1x, s1y, s1w, s1h, UI_WHITE, UI_BORDER, 1)
    # close btn
    add_rect(sl, s1x + s1w - Inches(0.38), s1y + Inches(0.1), Inches(0.28), Inches(0.28),
             RGBColor(0xEE, 0xEE, 0xEE))
    txbox(sl, "✕", s1x + s1w - Inches(0.38), s1y + Inches(0.1), Inches(0.28), Inches(0.28),
          size=9, color=UI_MUTED, align=PP_ALIGN.CENTER)
    # Icon placeholder
    add_rect(sl, s1x + Inches(1.55), s1y + Inches(0.55), Inches(0.6), Inches(0.6),
             RGBColor(0xD4, 0xED, 0xE4))
    txbox(sl, "✦", s1x + Inches(1.55), s1y + Inches(0.55), Inches(0.6), Inches(0.6),
          size=18, color=UI_GREEN, align=PP_ALIGN.CENTER)
    # Heading
    txbox(sl, "Find Your\nPerfect Match", s1x + Inches(0.3), s1y + Inches(1.35),
          s1w - Inches(0.6), Inches(0.9), size=17, bold=True, color=UI_HEADER,
          align=PP_ALIGN.CENTER)
    # Subheading
    txbox(sl, "Answer a few quick questions and we'll recommend the best products for you.",
          s1x + Inches(0.35), s1y + Inches(2.38), s1w - Inches(0.7), Inches(0.7),
          size=10, color=UI_MUTED, align=PP_ALIGN.CENTER, wrap=True)
    # CTA
    add_rect(sl, s1x + Inches(0.55), s1y + Inches(3.38), s1w - Inches(1.1), Inches(0.42),
             UI_ACCENT)
    txbox(sl, "Take the Quiz", s1x + Inches(0.55), s1y + Inches(3.38),
          s1w - Inches(1.1), Inches(0.42), size=11, bold=True,
          color=UI_WHITE, align=PP_ALIGN.CENTER)
    txbox(sl, "Takes ~1 minute", s1x + Inches(0.3), s1y + Inches(3.95),
          s1w - Inches(0.6), Inches(0.2), size=8, color=UI_MUTED, align=PP_ALIGN.CENTER)
    txbox(sl, "Step 1 — Intro", s1x, s1y + s1h + Inches(0.06),
          s1w, Inches(0.2), size=8, color=ACCENT, align=PP_ALIGN.CENTER, italic=True)

    # ── STATE 2: Question Screen ──────────────────────────────────────────
    s2x = Inches(4.82)
    s2y = s1y
    s2w = Inches(3.7)
    s2h = s1h
    add_rect_border(sl, s2x, s2y, s2w, s2h, UI_WHITE, UI_BORDER, 1)
    # close btn
    add_rect(sl, s2x + s2w - Inches(0.38), s2y + Inches(0.1), Inches(0.28), Inches(0.28),
             RGBColor(0xEE, 0xEE, 0xEE))
    txbox(sl, "✕", s2x + s2w - Inches(0.38), s2y + Inches(0.1), Inches(0.28), Inches(0.28),
          size=9, color=UI_MUTED, align=PP_ALIGN.CENTER)
    # Progress bar bg
    add_rect(sl, s2x + Inches(0.25), s2y + Inches(0.52), s2w - Inches(0.5), Inches(0.07),
             RGBColor(0xE5, 0xE5, 0xE5))
    # Progress bar fill (40% = 2 of 5)
    add_rect(sl, s2x + Inches(0.25), s2y + Inches(0.52),
             (s2w - Inches(0.5)) * 0.4, Inches(0.07), RGBColor(0x1A, 0x1A, 0x1A))
    txbox(sl, "QUESTION 2 OF 5", s2x + Inches(0.25), s2y + Inches(0.68),
          s2w - Inches(0.5), Inches(0.2), size=7, color=UI_MUTED)
    txbox(sl, "Which finish do you prefer?",
          s2x + Inches(0.25), s2y + Inches(0.94), s2w - Inches(0.5), Inches(0.45),
          size=12, bold=True, color=UI_HEADER, wrap=True)
    txbox(sl, "Select one", s2x + Inches(0.25), s2y + Inches(1.42),
          s2w - Inches(0.5), Inches(0.2), size=8, color=UI_MUTED, italic=True)
    # Options
    quiz_opts = ["Matte", "Glossy", "Satin", "Dewy"]
    for oi, opt in enumerate(quiz_opts):
        oy2 = s2y + Inches(1.68) + oi * Inches(0.58)
        selected = (oi == 0)
        obg = RGBColor(0x1A, 0x1A, 0x1A) if selected else UI_WHITE
        oborder = RGBColor(0x1A, 0x1A, 0x1A)
        add_rect_border(sl, s2x + Inches(0.25), oy2, s2w - Inches(0.5), Inches(0.44),
                        obg, oborder, 2)
        txbox(sl, opt, s2x + Inches(0.45), oy2 + Inches(0.1), s2w - Inches(0.7), Inches(0.24),
              size=10, color=UI_WHITE if selected else UI_TEXT, bold=selected)
    # Actions
    act_y = s2y + s2h - Inches(0.62)
    add_rect_border(sl, s2x + Inches(0.25), act_y, Inches(0.9), Inches(0.36),
                    UI_WHITE, UI_BORDER, 2)
    txbox(sl, "← Back", s2x + Inches(0.25), act_y, Inches(0.9), Inches(0.36),
          size=9, color=UI_TEXT, align=PP_ALIGN.CENTER)
    add_rect(sl, s2x + s2w - Inches(1.15), act_y, Inches(0.9), Inches(0.36),
             RGBColor(0x1A, 0x1A, 0x1A))
    txbox(sl, "Next →", s2x + s2w - Inches(1.15), act_y, Inches(0.9), Inches(0.36),
          size=9, bold=True, color=UI_WHITE, align=PP_ALIGN.CENTER)
    txbox(sl, "Step 2 — Question (single-select)", s2x, s2y + s2h + Inches(0.06),
          s2w, Inches(0.2), size=8, color=ACCENT, align=PP_ALIGN.CENTER, italic=True)

    # ── STATE 3: Results Screen ───────────────────────────────────────────
    s3x = Inches(9.08)
    s3y = s1y
    s3w = Inches(3.95)
    s3h = s1h
    add_rect_border(sl, s3x, s3y, s3w, s3h, UI_WHITE, UI_BORDER, 1)
    # close btn
    add_rect(sl, s3x + s3w - Inches(0.38), s3y + Inches(0.1), Inches(0.28), Inches(0.28),
             RGBColor(0xEE, 0xEE, 0xEE))
    txbox(sl, "✕", s3x + s3w - Inches(0.38), s3y + Inches(0.1), Inches(0.28), Inches(0.28),
          size=9, color=UI_MUTED, align=PP_ALIGN.CENTER)
    # Results header
    txbox(sl, "Your Recommendations", s3x + Inches(0.25), s3y + Inches(0.18),
          Inches(2.5), Inches(0.32), size=13, bold=True, color=UI_HEADER)
    # nav arrows
    add_rect_border(sl, s3x + s3w - Inches(0.78), s3y + Inches(0.22),
                    Inches(0.28), Inches(0.28), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "‹", s3x + s3w - Inches(0.78), s3y + Inches(0.22),
          Inches(0.28), Inches(0.28), size=10, color=UI_TEXT, align=PP_ALIGN.CENTER)
    add_rect_border(sl, s3x + s3w - Inches(0.45), s3y + Inches(0.22),
                    Inches(0.28), Inches(0.28), UI_WHITE, UI_BORDER, 1)
    txbox(sl, "›", s3x + s3w - Inches(0.45), s3y + Inches(0.22),
          Inches(0.28), Inches(0.28), size=10, color=UI_TEXT, align=PP_ALIGN.CENTER)
    # Product cards
    res_prods = [("Matte Fix\nFoundation", "$42"), ("Velvet Lip\nColour", "$28")]
    for pi2, (pname2, pprice) in enumerate(res_prods):
        pcx2 = s3x + Inches(0.22) + pi2 * Inches(1.82)
        pcy2 = s3y + Inches(0.62)
        # Image
        add_rect_border(sl, pcx2, pcy2, Inches(1.62), Inches(1.6),
                        UI_BG, UI_BORDER, 1)
        txbox(sl, "[ product\n  image ]", pcx2, pcy2, Inches(1.62), Inches(1.6),
              size=8, color=UI_MUTED, align=PP_ALIGN.CENTER)
        # Info
        txbox(sl, pname2, pcx2, pcy2 + Inches(1.65), Inches(1.62), Inches(0.38),
              size=8, bold=True, color=UI_TEXT, wrap=True)
        txbox(sl, pprice, pcx2, pcy2 + Inches(2.06), Inches(1.62), Inches(0.2),
              size=8, color=UI_MUTED)
        # ATC
        add_rect(sl, pcx2, pcy2 + Inches(2.28), Inches(1.62), Inches(0.32),
                 RGBColor(0x1A, 0x1A, 0x1A))
        txbox(sl, "Add to Cart", pcx2, pcy2 + Inches(2.28), Inches(1.62), Inches(0.32),
              size=8, bold=True, color=UI_WHITE, align=PP_ALIGN.CENTER)
    # retake
    txbox(sl, "retake quiz", s3x + Inches(0.25), s3y + s3h - Inches(0.5),
          s3w - Inches(0.5), Inches(0.28), size=9, color=UI_MUTED,
          align=PP_ALIGN.CENTER, italic=True)
    txbox(sl, "Step 3 — Results", s3x, s3y + s3h + Inches(0.06),
          s3w, Inches(0.2), size=8, color=ACCENT, align=PP_ALIGN.CENTER, italic=True)

    # Arrows between states
    for ax3, ay3 in [(Inches(4.32), Inches(4.15)), (Inches(8.63), Inches(4.15))]:
        txbox(sl, "→", ax3, ay3, Inches(0.5), Inches(0.4),
              size=18, bold=True, color=ACCENT2, align=PP_ALIGN.CENTER)

    slide_number(sl, n, total)

slides_data.append(slide_customer_mockup)


# ── 7. Key Features — Merchant Admin ─────────────────────────────────────────
def slide_merchant_features(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Merchant Admin — Features", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    features = [
        ("Campaign Manager",
         ["Create UTM-targeted campaigns",
          "Match by utm_source, medium, campaign",
          "AND / OR match mode",
          "Activate / deactivate anytime"]),
        ("Quiz Builder",
         ["Multi-step questionnaire editor",
          "Single-select & multi-select questions",
          "Drag-to-reorder questions",
          "Conditional branching between questions"]),
        ("Scoring Rules",
         ["Map each answer option to product tags",
          "Assign weight 1–10 per rule",
          "Support product ID or collection targeting",
          "Multiple rules per option accumulate"]),
        ("Preview & Test",
         ["Walk through quiz as a customer",
          "See ranked recommendations live",
          "Validate before going live",
          "Full scoring result with scores shown"]),
    ]
    col_w = Inches(3.0)
    gap = Inches(0.18)
    start_x = Inches(0.4)
    card_y = Inches(1.6)
    card_h = Inches(5.0)
    for i, (title, bullets) in enumerate(features):
        x = start_x + i * (col_w + gap)
        add_rect(sl, x, card_y, col_w, card_h, CARD_BG)
        add_rect(sl, x, card_y, col_w, Inches(0.07), ACCENT)
        txbox(sl, title, x + Inches(0.18), card_y + Inches(0.2),
              col_w - Inches(0.3), Inches(0.5), size=15, bold=True, color=ACCENT)
        for j, b in enumerate(bullets):
            by = card_y + Inches(0.85) + j * Inches(0.88)
            add_rect(sl, x + Inches(0.18), by - Inches(0.02), Inches(0.08), Inches(0.32), ACCENT)
            txbox(sl, b, x + Inches(0.35), by - Inches(0.05),
                  col_w - Inches(0.5), Inches(0.75), size=12, color=WHITE, wrap=True)
    slide_number(sl, n, total)

slides_data.append(slide_merchant_features)


# ── 8. Key Features — Customer-Facing ─────────────────────────────────────────
def slide_customer_features(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Customer-Facing — Quiz & Recommendations", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    left_items = [
        ("Inline Quiz Rendering",
         "Served directly on Shopify store pages via App Proxy — no pop-ups, no redirects, fully on-brand."),
        ("UTM-Aware Quiz Matching",
         "Only shows the quiz relevant to the visitor's campaign; different UTMs trigger different quizzes."),
        ("Dynamic Question Branching",
         "Conditional logic routes visitors to different follow-up questions based on their selection."),
        ("Progress Indicator",
         "Shows Q1 of 5 style counter and question type hint (single / multi-select) for better UX."),
    ]
    right_items = [
        ("Weighted Scoring Engine",
         "Each answer contributes a weight to product tags. Scores stack across all answers before ranking."),
        ("Real-Time Shopify Product Query",
         "After quiz completion, the engine queries Shopify GraphQL with the top-weighted tags."),
        ("Re-Ranked Results",
         "Products are re-ranked by their cumulative tag-weight score — most relevant product comes first."),
        ("Rich Product Cards",
         "Recommendations show image, title, price — with add-to-cart via Keen Slider carousel."),
    ]
    lx = Inches(0.5)
    rx = Inches(6.85)
    col_w = Inches(6.2)
    item_h = Inches(1.5)
    start_y = Inches(1.6)
    for col_x, items in [(lx, left_items), (rx, right_items)]:
        for i, (title, desc) in enumerate(items):
            y = start_y + i * item_h
            add_rect(sl, col_x, y, Inches(0.06), Inches(0.85), ACCENT)
            txbox(sl, title, col_x + Inches(0.18), y, col_w - Inches(0.2), Inches(0.38),
                  size=14, bold=True, color=WHITE)
            txbox(sl, desc, col_x + Inches(0.18), y + Inches(0.38),
                  col_w - Inches(0.2), Inches(0.85), size=12, color=LIGHT_GREY, wrap=True)
    add_rect(sl, Inches(6.7), Inches(1.5), Inches(0.03), Inches(5.5), CARD_BG)
    slide_number(sl, n, total)

slides_data.append(slide_customer_features)


# ── 9. Architecture Overview ──────────────────────────────────────────────────
def slide_architecture(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Architecture Overview", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    cols = [
        {"title": "Admin Frontend", "color": ACCENT,
         "items": ["React 18 + Shopify Polaris UI", "Raviger client-side routing",
                   "Shopify App Bridge (embedded)", "Campaign & Quiz Builder pages", "Vite build tooling"]},
        {"title": "Express Backend", "color": ACCENT2,
         "items": ["Node.js + Express 5", "JWT auth middleware (admin)",
                   "HMAC verify middleware (proxy)", "MongoDB / Mongoose models",
                   "Shopify Admin GraphQL client", "GDPR webhook handlers"]},
        {"title": "Customer Proxy", "color": RGBColor(0x6C, 0x8E, 0xFF),
         "items": ["App Proxy — no iframe needed", "UTM-campaign matching",
                   "Scoring & ranking engine", "Shopify product query (tags)",
                   "Session token tracking"]},
    ]
    col_w = Inches(3.6)
    gap   = Inches(0.42)
    start_x = Inches(0.6)
    cy = Inches(1.6)
    ch = Inches(5.3)
    for i, col in enumerate(cols):
        x = start_x + i * (col_w + gap)
        add_rect(sl, x, cy, col_w, ch, CARD_BG)
        add_rect(sl, x, cy, col_w, Inches(0.08), col["color"])
        txbox(sl, col["title"], x + Inches(0.2), cy + Inches(0.18),
              col_w - Inches(0.35), Inches(0.5), size=16, bold=True, color=col["color"])
        for j, item in enumerate(col["items"]):
            iy = cy + Inches(0.85) + j * Inches(0.8)
            add_rect(sl, x + Inches(0.2), iy + Inches(0.12), Inches(0.12), Inches(0.12), col["color"])
            txbox(sl, item, x + Inches(0.42), iy, col_w - Inches(0.55), Inches(0.65),
                  size=12, color=WHITE, wrap=True)
        if i < len(cols) - 1:
            ax = x + col_w + Inches(0.05)
            txbox(sl, "⇄", ax, cy + Inches(2.5), gap, Inches(0.5),
                  size=20, bold=True, color=LIGHT_GREY, align=PP_ALIGN.CENTER)
    txbox(sl, "Shared Data Store: MongoDB (Campaigns · Questions · Options · Scoring Rules · Quiz Sessions · Store)",
          Inches(0.6), Inches(6.85), Inches(12.1), Inches(0.4),
          size=12, color=ACCENT, align=PP_ALIGN.CENTER)
    slide_number(sl, n, total)

slides_data.append(slide_architecture)


# ── 10. Data Model ─────────────────────────────────────────────────────────────
def slide_data_model(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "Data Model", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    entities = [
        ("Campaign", ACCENT,
         ["name", "utm_source / medium / campaign", "match_mode (all|any)", "is_active", "result_tag_limit", "result_product_limit"]),
        ("Question", RGBColor(0x6C, 0x8E, 0xFF),
         ["campaign_id", "text", "type (single|multi)", "order", "is_required"]),
        ("Option", ACCENT2,
         ["question_id", "label", "order", "next_question_id"]),
        ("ScoringRule", RGBColor(0xFF, 0x7F, 0x50),
         ["option_id", "target_type", "target_value (tag/id)", "weight (1–10)"]),
        ("QuizSession", RGBColor(0x9B, 0x59, 0xB6),
         ["campaign_id, shop", "answers[ ]", "result_product_ids", "utm params", "completed_at"]),
    ]
    ew = Inches(2.3)
    eg = Inches(0.15)
    sx = Inches(0.4)
    ey = Inches(1.55)
    eh = Inches(5.2)
    for i, (name, color, fields) in enumerate(entities):
        x = sx + i * (ew + eg)
        add_rect(sl, x, ey, ew, eh, CARD_BG)
        add_rect(sl, x, ey, ew, Inches(0.5), color)
        txbox(sl, name, x + Inches(0.1), ey + Inches(0.07),
              ew - Inches(0.15), Inches(0.38), size=14, bold=True,
              color=DARK_BG, align=PP_ALIGN.CENTER)
        for j, f in enumerate(fields):
            fy = ey + Inches(0.65) + j * Inches(0.72)
            add_rect(sl, x + Inches(0.15), fy + Inches(0.14), Inches(0.1), Inches(0.1), color)
            txbox(sl, f, x + Inches(0.35), fy, ew - Inches(0.45), Inches(0.65),
                  size=10, color=WHITE, wrap=True)
    txbox(sl, "Campaign  1:N  Question  1:N  Option  1:N  ScoringRule",
          Inches(0.4), Inches(7.0), Inches(9), Inches(0.35),
          size=11, color=ACCENT, italic=True)
    txbox(sl, "QuizSession  N:1  Campaign",
          Inches(9.5), Inches(7.0), Inches(3.5), Inches(0.35),
          size=11, color=ACCENT, italic=True, align=PP_ALIGN.RIGHT)
    slide_number(sl, n, total)

slides_data.append(slide_data_model)


# ── 11. End-to-End Merchant Journey ────────────────────────────────────────────
def slide_merchant_journey(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "End-to-End — Merchant Journey", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    steps = [
        ("1  Create Campaign",
         ["Name the campaign", "Set utm_source / utm_medium / utm_campaign",
          "Choose AND or OR match mode", "Set tag and product result limits"]),
        ("2  Build Quiz",
         ["Add questions (single or multi-select)", "Add 2+ answer options per question",
          "Reorder questions with up/down arrows", "Optional: link options to jump questions"]),
        ("3  Define Scoring Rules",
         ["For each option, add scoring rules", "Map option → product tag + weight (1–10)",
          "Multiple rules per option allowed", "Higher weight = stronger signal"]),
        ("4  Preview & Activate",
         ["Click Preview to walk through quiz", "Verify ranked recommendations are correct",
          "Toggle campaign to Active", "Share UTM link in campaigns"]),
    ]
    bw = Inches(2.9)
    gap = Inches(0.27)
    sx = Inches(0.5)
    by_start = Inches(1.55)
    bh = Inches(5.2)
    for i, (title, items) in enumerate(steps):
        x = sx + i * (bw + gap)
        add_rect(sl, x, by_start, bw, bh, CARD_BG)
        add_rect(sl, x, by_start, bw, Inches(0.07), ACCENT)
        txbox(sl, title, x + Inches(0.18), by_start + Inches(0.15),
              bw - Inches(0.3), Inches(0.5), size=14, bold=True, color=ACCENT)
        for j, item in enumerate(items):
            iy = by_start + Inches(0.85) + j * Inches(0.95)
            add_rect(sl, x + Inches(0.2), iy + Inches(0.14), Inches(0.1), Inches(0.1), ACCENT)
            txbox(sl, item, x + Inches(0.4), iy, bw - Inches(0.5), Inches(0.8),
                  size=12, color=WHITE, wrap=True)
        if i < len(steps) - 1:
            ax = x + bw + Inches(0.03)
            txbox(sl, "→", ax, by_start + Inches(2.3), gap + Inches(0.1), Inches(0.45),
                  size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    slide_number(sl, n, total)

slides_data.append(slide_merchant_journey)


# ── 12. End-to-End Customer Journey ────────────────────────────────────────────
def slide_customer_journey(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "End-to-End — Customer Journey", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    steps = [
        ("Visit via UTM Link", RGBColor(0x00, 0xB5, 0x8A),
         "Customer clicks an Instagram / email / ad link carrying utm_source, utm_medium, and utm_campaign query params."),
        ("Quiz Loads", RGBColor(0x6C, 0x8E, 0xFF),
         "App Proxy matches UTM params to an active campaign, renders the first question inline on the Shopify page."),
        ("Answer Questions", ACCENT2,
         "Customer picks single or multi-select answers. Conditional branching routes to the next relevant question. Progress bar shown."),
        ("Submit Quiz", RGBColor(0xFF, 0x7F, 0x50),
         "After the final question, customer clicks 'See My Results'. Selected option IDs sent to the scoring endpoint."),
        ("Scoring Engine", RGBColor(0x9B, 0x59, 0xB6),
         "Server aggregates weights per tag from all selected options, sorts tags by total score, queries Shopify products by top tags."),
        ("View Results", RGBColor(0xE7, 0x4C, 0x3C),
         "Ranked product cards displayed in a carousel with image, title, price. Customer adds to cart directly from the quiz modal."),
    ]
    step_h = Inches(0.92)
    sx = Inches(0.6)
    sy = Inches(1.55)
    for i, (title, color, desc) in enumerate(steps):
        y = sy + i * (step_h + Inches(0.06))
        add_rect(sl, sx, y, Inches(12.0), step_h, CARD_BG)
        add_rect(sl, sx, y, Inches(0.12), step_h, color)
        txbox(sl, str(i + 1), sx + Inches(0.2), y + Inches(0.2),
              Inches(0.4), Inches(0.5), size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
        txbox(sl, title, sx + Inches(0.7), y + Inches(0.1), Inches(2.2), Inches(0.42),
              size=13, bold=True, color=color)
        txbox(sl, desc, sx + Inches(0.7), y + Inches(0.5), Inches(11.1), Inches(0.38),
              size=11, color=LIGHT_GREY, wrap=True)
    slide_number(sl, n, total)

slides_data.append(slide_customer_journey)


# ── 13. Scoring Algorithm ─────────────────────────────────────────────────────
def slide_scoring(prs, n, total):
    sl = prs.slides.add_slide(BLANK)
    bg(sl)
    accent_bar(sl)
    txbox(sl, "How the Scoring Algorithm Works", Inches(0.6), Inches(0.7), Inches(12), Inches(0.6),
          size=32, bold=True, color=WHITE)

    add_rect(sl, Inches(0.5), Inches(1.55), Inches(5.8), Inches(5.4), CARD_BG)
    add_rect(sl, Inches(0.5), Inches(1.55), Inches(5.8), Inches(0.07), ACCENT)
    txbox(sl, "Step-by-Step", Inches(0.68), Inches(1.68), Inches(5.5), Inches(0.45),
          size=15, bold=True, color=ACCENT)

    algo_steps = [
        ("Collect",      "Gather all option_ids the customer selected across every question."),
        ("Fetch Rules",  "Query ScoringRule for those option_ids → list of {target_value, weight}."),
        ("Aggregate",    "Sum weights per tag: scores['matte'] += weight for each matching rule."),
        ("Sort Tags",    "Sort tags by total score descending, take top N (result_tag_limit)."),
        ("Query Shopify","GraphQL: products(query: 'tag:matte OR tag:glossy OR …') from catalog."),
        ("Re-Rank",      "Each product's score = sum of scores for its tags that appear in our map."),
        ("Slice",        "Return top M products (result_product_limit) sorted by score."),
    ]
    for i, (step, desc) in enumerate(algo_steps):
        y = Inches(2.25) + i * Inches(0.65)
        add_rect(sl, Inches(0.65), y + Inches(0.1), Inches(0.08), Inches(0.28), ACCENT)
        txbox(sl, step, Inches(0.82), y, Inches(1.1), Inches(0.38), size=12, bold=True, color=ACCENT)
        txbox(sl, desc, Inches(2.0), y, Inches(4.1), Inches(0.55),
              size=11, color=LIGHT_GREY, wrap=True)

    rx = Inches(6.65)
    rw = Inches(6.2)
    add_rect(sl, rx, Inches(1.55), rw, Inches(5.4), CARD_BG)
    add_rect(sl, rx, Inches(1.55), rw, Inches(0.07), ACCENT2)
    txbox(sl, "Example Walkthrough", rx + Inches(0.2), Inches(1.68),
          rw - Inches(0.3), Inches(0.45), size=15, bold=True, color=ACCENT2)

    example = [
        ("Customer answers:", WHITE, False),
        ("  Q1 → 'Matte finish'  (tag: matte, w=8)  (tag: long-lasting, w=5)", LIGHT_GREY, False),
        ("  Q2 → 'Daytime look'  (tag: natural, w=7)  (tag: matte, w=4)", LIGHT_GREY, False),
        ("", WHITE, False),
        ("Aggregated scores:", WHITE, False),
        ("  matte        = 8 + 4 = 12  ← highest", ACCENT, True),
        ("  long-lasting = 5", LIGHT_GREY, False),
        ("  natural      = 7", LIGHT_GREY, False),
        ("", WHITE, False),
        ("Top tags queried in Shopify: matte, natural, long-lasting", LIGHT_GREY, False),
        ("", WHITE, False),
        ("Product re-ranking:", WHITE, False),
        ("  P1 [matte + natural]:      12+7 = 19  ← #1", ACCENT, True),
        ("  P3 [natural + long-lasting]: 7+5 = 12  ← #2", LIGHT_GREY, False),
        ("  P2 [matte]:                   12 = 12  ← #2", LIGHT_GREY, False),
    ]
    ty = Inches(2.2)
    for line, col, bold in example:
        if line == "":
            ty += Inches(0.08)
            continue
        txbox(sl, line, rx + Inches(0.2), ty, rw - Inches(0.3), Inches(0.32),
              size=10, color=col, bold=bold)
        ty += Inches(0.33)

    slide_number(sl, n, total)

slides_data.append(slide_scoring)


# ══════════════════════════════════════════════════════════════════════════════
# Render
# ══════════════════════════════════════════════════════════════════════════════

for idx, fn in enumerate(slides_data, start=1):
    fn(prs, idx, TOTAL)

out = "/Users/anit/Desktop/anit/apps/recommendation-engine/Recommendation_Engine_Presentation.pptx"
prs.save(out)
print(f"Saved → {out}  ({TOTAL} slides)")
